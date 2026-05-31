#!/usr/bin/env python3
"""
无人机城市交通态势感知系统 — CTO技术汇报PPT生成器 v3
Deep navy theme with cyan/teal accents
Enhanced: more visual elements, better layout variety, deeper technical content
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import math

# ──────────────────── Theme Colors ────────────────────
BG_DARK    = RGBColor(0x0B, 0x14, 0x26)   # Deep navy
BG_CARD    = RGBColor(0x11, 0x1E, 0x36)   # Card background
BG_CARD2   = RGBColor(0x16, 0x25, 0x40)   # Slightly lighter card
BG_CARD3   = RGBColor(0x1A, 0x2B, 0x45)   # Even lighter
ACCENT     = RGBColor(0x00, 0xD4, 0xAA)   # Cyan/teal
ACCENT2    = RGBColor(0x38, 0x8B, 0xFD)   # Blue
ACCENT3    = RGBColor(0xFF, 0x6B, 0x35)   # Orange
ACCENT4    = RGBColor(0xA8, 0x55, 0xF7)   # Purple
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xB0, 0xBC, 0xD0)
MID_GRAY   = RGBColor(0x7A, 0x8A, 0xA0)
DARK_LINE   = RGBColor(0x20, 0x30, 0x50)
RED         = RGBColor(0xEF, 0x44, 0x44)
YELLOW      = RGBColor(0xFB, 0xBF, 0x24)
GREEN       = RGBColor(0x22, 0xC5, 0x5E)
DARK_GREEN  = RGBColor(0x0A, 0x2E, 0x1A)
DARK_BLUE   = RGBColor(0x0A, 0x14, 0x2E)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height

# ──────────────────── Helper Functions ────────────────────
def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, fill_color=None, border_color=None, border_width=Pt(0), shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.line.width = border_width
    if border_color:
        shape.line.color.rgb = border_color
    else:
        shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    return shape

def add_rect(slide, left, top, width, height, fill_color=None, border_color=None, border_width=Pt(0)):
    return add_shape(slide, left, top, width, height, fill_color, border_color, border_width, MSO_SHAPE.RECTANGLE)

def add_text_box(slide, left, top, width, height, text, font_size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_multiline_text(slide, left, top, width, height, lines, font_size=13, color=WHITE, line_spacing=1.2, font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, str):
            text, c, b, s = item, color, False, font_size
        elif len(item) == 2:
            text, c = item; b, s = False, font_size
        elif len(item) == 3:
            text, c, b = item; s = font_size
        else:
            text, c, b, s = item
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(s)
        p.font.color.rgb = c
        p.font.bold = b
        p.font.name = font_name
        p.space_after = Pt(s * 0.3)
    return txBox

def add_card(slide, left, top, width, height, title, items, accent_color=ACCENT, icon_text=""):
    card = add_shape(slide, left, top, width, height, fill_color=BG_CARD, border_color=DARK_LINE, border_width=Pt(1))
    add_rect(slide, left + Inches(0.15), top + Inches(0.1), Inches(0.4), Pt(4), fill_color=accent_color)
    add_text_box(slide, left + Inches(0.5), top + Inches(0.05), width - Inches(0.6), Inches(0.45), title, font_size=15, color=WHITE, bold=True)
    lines = []
    for item in items:
        if isinstance(item, tuple):
            lines.append(item)
        else:
            lines.append((f"• {item}", LIGHT_GRAY, False, 12))
    add_multiline_text(slide, left + Inches(0.25), top + Inches(0.55), width - Inches(0.5), height - Inches(0.65), lines, font_size=12)

def add_flow_box(slide, left, top, width, height, text, fill_color=BG_CARD2, text_color=WHITE, font_size=11, border_color=None):
    bc = border_color or ACCENT
    shape = add_shape(slide, left, top, width, height, fill_color=fill_color, border_color=bc, border_width=Pt(1.5))
    shape.text_frame.word_wrap = True
    p = shape.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = text_color
    p.font.bold = True
    p.font.name = "Microsoft YaHei"
    p.alignment = PP_ALIGN.CENTER
    shape.text_frame.paragraphs[0].space_before = Pt(0)
    return shape

def add_arrow(slide, left, top, width, height, color=ACCENT):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_down_arrow(slide, left, top, width, height, color=ACCENT):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_page_number(slide, num, total):
    add_text_box(slide, Inches(12.3), Inches(7.1), Inches(0.8), Inches(0.3), f"{num}/{total}", font_size=10, color=MID_GRAY, alignment=PP_ALIGN.RIGHT)

def add_section_header(slide, title, subtitle=""):
    add_rect(slide, Inches(0), Inches(0), W, Pt(4), fill_color=ACCENT)
    add_text_box(slide, Inches(0.8), Inches(0.35), Inches(11), Inches(0.7), title, font_size=28, color=WHITE, bold=True)
    if subtitle:
        add_text_box(slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.4), subtitle, font_size=14, color=LIGHT_GRAY)

def add_kpi_box(slide, left, top, width, height, value, label, color=ACCENT):
    add_shape(slide, left, top, width, height, fill_color=BG_CARD, border_color=color, border_width=Pt(2))
    add_text_box(slide, left, top + Inches(0.15), width, Inches(0.5), value, font_size=26, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, left, top + Inches(0.6), width, Inches(0.3), label, font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

def add_icon_circle(slide, left, top, size, text, color=ACCENT):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    p = shape.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    return shape

TOTAL_SLIDES = 20

# ════════════════════════════════════════════════════════
# SLIDE 1: Cover
# ════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9), Inches(-1), Inches(6), Inches(6))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0x0F, 0x1C, 0x33)
shape.line.fill.background()

shape2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2), Inches(4.5), Inches(5), Inches(5))
shape2.fill.solid()
shape2.fill.fore_color.rgb = RGBColor(0x0D, 0x18, 0x2E)
shape2.line.fill.background()

for i in range(6):
    x = Inches(10 + i * 0.4)
    y = Inches(0.5 + i * 0.3)
    s = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, x, y, Inches(0.35), Inches(0.35))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor(0x15, 0x25, 0x40)
    s.line.color.rgb = ACCENT
    s.line.width = Pt(0.5)

add_rect(slide, Inches(1.2), Inches(2.5), Inches(1.0), Pt(5), fill_color=ACCENT)

add_text_box(slide, Inches(1.2), Inches(2.8), Inches(10), Inches(1.2),
    "无人机城市交通态势感知系统", font_size=44, color=WHITE, bold=True)
add_text_box(slide, Inches(1.2), Inches(4.0), Inches(10), Inches(0.6),
    "UAV-Based Urban Traffic Situation Awareness System", font_size=18, color=ACCENT)
add_text_box(slide, Inches(1.2), Inches(4.8), Inches(10), Inches(0.5),
    "技术方案汇报  |  CTO技术评审  |  2026.05", font_size=16, color=LIGHT_GRAY)

add_rect(slide, Inches(0), Inches(6.7), W, Inches(0.8), fill_color=RGBColor(0x08, 0x10, 0x1E))
add_text_box(slide, Inches(1.2), Inches(6.85), Inches(5), Inches(0.4),
    "YOLO11 + ByteTrack  |  Kafka + InfluxDB  |  双模式作业  |  49/49 E2E测试通过", font_size=13, color=MID_GRAY)
add_text_box(slide, Inches(9), Inches(6.85), Inches(3.5), Inches(0.4),
    "2026.05.31", font_size=13, color=MID_GRAY, alignment=PP_ALIGN.RIGHT)

print("Slide 1: Cover done")
