from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "generated"
SCREEN_DIR = OUT_DIR / "screenshots"
DOCX_PATH = OUT_DIR / "TrafficAnalyzer_路口车道标注功能详细技术方案.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(107, 114, 128)
LIGHT_FILL = "F2F4F7"
CALL_FILL = "EEF6F3"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def format_cell(cell, bold: bool = False, color: RGBColor | None = None, size: int = 9) -> None:
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.name = "Calibri"
            r._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
            r.font.size = Pt(size)
            r.font.bold = bold
            if color:
                r.font.color.rgb = color


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        set_cell_shading(hdr[i], LIGHT_FILL)
        format_cell(hdr[i], bold=True, color=INK, size=9)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            format_cell(cells[i], size=9)
    doc.add_paragraph()


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9120])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALL_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.bold = True
    run.font.color.rgb = DARK_BLUE
    run.font.size = Pt(10)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    for r in p2.runs:
        r.font.size = Pt(10)
        r.font.color.rgb = INK
    doc.add_paragraph()


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    r.italic = True


def add_screenshot(doc: Document, filename: str, caption: str) -> None:
    path = SCREEN_DIR / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.25))
    add_caption(doc, caption)


def add_code_block(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9120])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for line in text.splitlines():
        if p.text:
            p.add_run("\n")
        r = p.add_run(line)
        r.font.name = "Courier New"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(17, 24, 39)
    doc.add_paragraph()


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "TrafficAnalyzer 路口车道标注功能技术方案"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = MUTED


def title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("TrafficAnalyzer")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = DARK_BLUE

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("路口车道标注功能详细技术方案")
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = RGBColor(11, 37, 69)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("覆盖无人机悬停识别、自动生成标注任务、人工车道标注、参数持久化与后续检测复用")
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["产品定位", "无人机路口交通态势监测平台，面向巡飞视频检测、交通流统计、告警和 GIS 指挥。"],
            ["本方案范围", "标定中心 -> 车道标注；监控流程检测器；Kafka/Platform/Console 联动。"],
            ["生成日期", datetime.now().strftime("%Y-%m-%d")],
            ["当前验证数据", "inter_xqh 视频 + SRT 遥测，标注任务 lane-INT_camera_3-1781271962。"],
        ],
        [2200, 6920],
    )

    add_callout(
        doc,
        "方案结论",
        "系统采用“检测管道发现悬停 -> Kafka 发送快照与遥测 -> 平台生成标注任务 -> Console 人工绘制车道 -> 导出 roads/lanes JSON -> 后续 PipelineManager 自动复用”的闭环。该路径保持检测流、平台管理和前端标定页面解耦，并允许离线视频回放与真实无人机 RTSP 任务使用同一套契约。",
    )
    doc.add_page_break()


def build_doc() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    title_page(doc)

    doc.add_heading("1. 产品与建设目标", level=1)
    doc.add_paragraph(
        "TrafficAnalyzer 是无人机路口交通态势监测平台，核心管道从 MP4、RTSP 或摄像头读取视频，完成 YOLO 检测、ByteTrack 跟踪、无人机运动补偿、道路/车道统计，并通过 Kafka、WebSocket 和 Console 页面形成实时监控闭环。本方案聚焦巡飞转悬停后的车道标注能力。"
    )
    doc.add_paragraph(
        "业务目标是减少每次任务启动前手工准备车道参数的成本。当无人机在某个路口稳定悬停超过阈值后，系统自动截取检测画面生成待标注任务；人工在标定中心完成车道多边形绘制并保存；后续同一路口再次启动检测流时，平台优先使用已保存的标注参数。"
    )
    add_screenshot(doc, "01-sign-in.png", "图 1：Console 登录入口，统一进入实时监测、视频分析、告警和标定模块。")

    doc.add_page_break()
    doc.add_heading("2. 功能范围", level=1)
    add_table(
        doc,
        ["能力", "说明", "验收要点"],
        [
            ["悬停识别", "检测管道根据 MotionCompensationNode 输出的 is_hovering 与 drone_position 判断稳定悬停。", "同一路口固定坐标超过 30 秒，且半径误差不超过配置值。"],
            ["任务生成", "KafkaProducerNode 在 stats 消息中带 road_polygons、drone_position、is_hovering 与 annotation_snapshot_jpeg。", "平台 Consumer 收到稳定悬停消息后创建 lane task，包含图片 URL 和尺寸。"],
            ["人工标注", "标定中心的车道标注页从任务列表读取图片，SVG 画布点击顶点、闭合多边形并保存。", "保存 payload 包含 lane_id、direction、polygon、roads。"],
            ["参数复用", "LaneAnnotationStore 导出路口级 JSON，PipelineManager 在启动检测流时解析并优先使用。", "后续同 intersection_id 未显式指定 roads_json 时自动使用 export_path。"],
        ],
        [1800, 4700, 2620],
    )

    doc.add_heading("3. 页面截图与操作流", level=1)
    add_screenshot(doc, "04-dashboard.png", "图 2：指挥大屏，展示活跃路口、总车流量、拥堵指数和系统状态。")
    add_screenshot(doc, "03-monitoring.png", "图 3：实时监测页面，操作员从路口/无人机选择检测流，并观察视频与 BEV 地图。")
    add_screenshot(doc, "02-calibration-lane-task.png", "图 4：标定中心车道标注页，左侧为悬停生成的任务，中间为真实无人机快照画布。")

    doc.add_heading("4. 总体架构", level=1)
    doc.add_paragraph("车道标注闭环跨越检测管道、Kafka 消息、平台服务和 React Console 四层。")
    add_code_block(
        doc,
        """无人机视频/离线视频
  -> main_optimized.py 三进程检测管道
  -> MotionCompensationNode 输出 drone_position / is_hovering
  -> KafkaProducerNode 输出 statistics_N + annotation_snapshot_jpeg
  -> Platform KafkaConsumerService.observe_stats()
  -> LaneAnnotationStore 创建 lane task 与 JPEG 图片
  -> Console 标定中心拉取任务并绘制车道
  -> Calibration API 保存 lanes/roads
  -> 导出 lane_annotations/{intersection_id}.json
  -> PipelineManager 后续启动同一路口检测流时自动复用""",
    )
    add_table(
        doc,
        ["层级", "组件", "责任"],
        [
            ["检测管道", "main_optimized.py / KafkaProducerNode", "计算交通指标、发送 stats、编码用于标注的 JPEG 快照。"],
            ["消息层", "Kafka statistics_N", "把无人机位置、悬停状态、路口多边形和图片快照传递到平台。"],
            ["平台层", "KafkaConsumerService / LaneAnnotationStore / Calibration API", "识别稳定悬停、生成任务、保存标注、暴露图片与参数接口。"],
            ["前端层", "traffic-fly-console 标定中心", "任务列表、车道绘制画布、方向选择、保存标注参数。"],
            ["复用层", "PipelineManager", "启动检测流时解析同一路口保存的 export_path。"],
        ],
        [1500, 3300, 4320],
    )

    doc.add_heading("5. 核心流程设计", level=1)
    doc.add_heading("5.1 监控流程启动", level=2)
    doc.add_paragraph(
        "操作员在实时监测页面选择路口后点击“启动流”。前端通过 Pipeline API 传入 drone_id、intersection_id、video_src、telemetry_source 和 telemetry_file_path。平台创建子进程运行 main_optimized.py，并通过环境变量注入 VIDEO_SRC、ROADS_JSON、TOPIC_NAME、CAMERA_ID、INTERSECTION_ID 和 VIDEO_PORT。"
    )
    doc.add_heading("5.2 悬停检测与快照生产", level=2)
    doc.add_paragraph(
        "MotionCompensationNode 基于遥测产生无人机 ENU 位置、速度和 is_hovering。KafkaProducerNode 周期性发送 stats，若当前帧存在 frame_result 或 frame，则按配置宽度压缩成 JPEG，并以 base64 字段 annotation_snapshot_jpeg 附在消息内。"
    )
    doc.add_heading("5.3 平台生成标注任务", level=2)
    doc.add_paragraph(
        "KafkaConsumerService 在 _handle_stats 中调用 LaneAnnotationStore.observe_stats。Store 按 intersection_id 维护悬停状态：非悬停清空状态；坐标变化超过 hover_radius_m 重新计时；稳定时间超过 hover_seconds 且该路口尚无 active annotation 时，创建 pending task，并把 JPEG 解码保存到 lane_task_images。"
    )
    doc.add_heading("5.4 人工标注与保存", level=2)
    doc.add_paragraph(
        "Console 标定中心通过 GET /api/v1/calibration/lane-tasks 拉取任务。用户在 SVG 画布上点击顶点，至少 3 个点后闭合为车道，选择方向 straight、left_turn、right_turn 或 u_turn。保存时提交 lanes 和 roads，平台规范化 polygon 数值并写入 DB。"
    )
    doc.add_heading("5.5 后续检测复用", level=2)
    doc.add_paragraph(
        "保存成功后，Store 导出 lane_annotations/{intersection_id}.json。Pipeline API 收到新的启动请求时，如果调用方未显式指定 roads_json，则根据 intersection_id 查找已保存 annotation，并把 export_path 作为实际 roads_json。检测管道中的 VideoReader 和 LaneDetectionNode 后续可直接加载 manual 标注。"
    )

    doc.add_heading("6. 数据契约", level=1)
    add_table(
        doc,
        ["消息字段", "类型", "说明"],
        [
            ["intersection_id", "string", "统一路口 ID，例如 INT_camera_3。"],
            ["drone_position", "object", "包含 anchor_lat、anchor_lon、easting_m、northing_m。"],
            ["is_hovering", "boolean", "检测管道判断是否悬停。"],
            ["road_polygons", "object", "当前检测配置中的道路多边形，用于保存标注时继承 roads。"],
            ["annotation_snapshot_jpeg", "string", "base64 JPEG；平台创建任务时解码保存为图片。"],
            ["annotation_snapshot_width/height", "number", "前端 SVG viewBox 和图片尺寸。"],
        ],
        [2500, 1300, 5320],
    )
    add_table(
        doc,
        ["REST API", "方法", "用途"],
        [
            ["/api/v1/calibration/lane-tasks", "GET", "列出悬停生成的车道标注任务。"],
            ["/api/v1/calibration/lane-tasks/{task_id}/image", "GET", "返回任务快照 JPEG；图片地址可被 SVG image 直接加载。"],
            ["/api/v1/calibration/lane-tasks/{task_id}/annotation", "POST", "保存人工标注 lanes 与 roads。"],
            ["/api/v1/calibration/lane-annotations", "GET", "列出已保存可复用的路口车道参数。"],
            ["/api/v1/pipelines", "POST", "启动检测流；默认使用同路口已保存的标注参数。"],
        ],
        [4700, 1000, 3420],
    )

    doc.add_heading("7. 存储与文件结构", level=1)
    add_code_block(
        doc,
        """output/inter_xqh_page_e2e/
  lane_annotation_db.json
  lane_task_images/
    lane-INT_camera_3-1781271962.jpg
  lane_annotations/
    INT_camera_3.json""",
    )
    doc.add_paragraph(
        "lane_annotation_db.json 是轻量任务库，包含 tasks 与 annotations 两个集合。图片按 task_id 命名，导出参数按 intersection_id 命名，便于后续 PipelineManager 无状态读取。"
    )

    doc.add_heading("8. 前端交互设计", level=1)
    doc.add_paragraph(
        "车道标注页采用三栏布局：左栏任务列表，中间 SVG 绘制画布，右栏当前标注摘要和保存按钮。画布以任务图片尺寸作为 viewBox，点击坐标直接落到原图坐标系，避免额外换算误差。"
    )
    add_table(
        doc,
        ["交互", "规则"],
        [
            ["选择任务", "默认选择最新任务，也可点击左侧任务卡切换。"],
            ["添加顶点", "点击画布添加点，点位以整数像素坐标记录。"],
            ["闭合车道", "至少 3 个顶点才能闭合为 lane polygon。"],
            ["方向选择", "下拉选择直行、左转、右转、掉头。"],
            ["保存", "只有存在 selectedTask 且至少一个 draft lane 时可保存。"],
        ],
        [2200, 6920],
    )

    doc.add_heading("9. 运行配置", level=1)
    add_table(
        doc,
        ["配置项", "默认/示例", "说明"],
        [
            ["LANE_ANNOTATION_DB_PATH", "output/.../lane_annotation_db.json", "任务库路径，决定图片和导出目录。"],
            ["LANE_ANNOTATION_HOVER_SECONDS", "30", "稳定悬停时长阈值。"],
            ["LANE_ANNOTATION_HOVER_RADIUS_M", "1.5", "同一悬停点的最大偏移半径。"],
            ["annotation_snapshot_width", "960", "KafkaProducerNode 快照压缩宽度。"],
            ["annotation_snapshot_jpeg_quality", "75", "JPEG 压缩质量。"],
        ],
        [3000, 2500, 3620],
    )

    doc.add_heading("10. 安全与权限", level=1)
    doc.add_paragraph(
        "平台 API 默认走 JWT 中间件。任务图片端点需要允许浏览器 SVG image 直接读取，因此 GET /lane-tasks/{task_id}/image 作为只读公共图片端点处理；保存标注的 POST 仍受 JWT 保护，避免未授权写入车道参数。"
    )
    add_callout(
        doc,
        "权限边界",
        "公共化的仅是任务图片读取，不包括任务列表、annotation 保存或参数库查询。生产环境如需更强隔离，可将图片 URL 改为短期签名地址或由前端以 Blob URL 加载。"
    )

    doc.add_heading("11. 验证方案", level=1)
    add_table(
        doc,
        ["验证项", "命令/方式", "预期"],
        [
            ["后端语法", "python -m py_compile 相关 Python 文件", "无语法错误。"],
            ["前端类型", "pnpm -C traffic-fly-console exec tsc -b", "TypeScript 编译通过。"],
            ["任务图片", "GET /api/v1/calibration/lane-tasks/{task_id}/image", "返回 image/jpeg，尺寸与任务字段一致。"],
            ["未授权写入", "无 token POST annotation", "返回 401。"],
            ["标注复用", "POST /api/v1/pipelines 不传 roads_json", "返回 roads_json 为该路口导出 JSON 路径。"],
            ["页面截图", "登录、Dashboard、实时监测、标定中心", "页面非空，标定页显示真实无人机快照。"],
        ],
        [1800, 4200, 3120],
    )

    doc.add_heading("12. 风险与演进", level=1)
    add_table(
        doc,
        ["风险", "影响", "建议"],
        [
            ["不同高度/角度重复悬停", "同一路口可能需要多套标注。", "后续引入 altitude/gimbal_pitch 分桶或版本化 annotation。"],
            ["图片公共读取", "内部网络可直接访问快照。", "生产环境改签名 URL 或鉴权后 Blob 加载。"],
            ["文件型 DB 并发", "多实例平台可能产生写冲突。", "迁移到 PostgreSQL 表并使用唯一索引约束 pending task。"],
            ["手绘精度", "人工误点会影响车道统计。", "加入撤销、拖拽顶点、已有车道编辑、缩放和平移。"],
            ["道路数量历史硬编码", "部分统计仍假定 road_1~road_5。", "继续推进动态道路数组在 Telegraf/Grafana 的兼容改造。"],
        ],
        [2100, 2900, 4120],
    )

    doc.add_page_break()
    doc.add_heading("13. 源码落点", level=1)
    add_table(
        doc,
        ["文件", "作用"],
        [
            ["nodes/KafkaProducerNode.py", "发送 stats、路口多边形、悬停快照和无人机位置。"],
            ["platform/app/kafka/consumer.py", "消费 stats 并调用 LaneAnnotationStore.observe_stats。"],
            ["platform/app/services/lane_annotation_store.py", "任务生成、图片保存、annotation 持久化和导出。"],
            ["platform/app/api/v1/calibration.py", "车道标注任务、图片、保存和参数库 API。"],
            ["platform/app/api/v1/pipelines.py", "启动检测流时复用已保存 annotation export_path。"],
            ["traffic-fly-console/src/features/calibration/index.tsx", "车道标注页面和 SVG 绘制逻辑。"],
            ["traffic-fly-console/src/features/monitoring/index.tsx", "检测流启动入口。"],
        ],
        [4300, 4820],
    )

    doc.add_heading("14. 交付结论", level=1)
    doc.add_paragraph(
        "该方案把无人机悬停检测、标注任务生产、人工车道绘制和检测参数复用串成闭环。它不改变既有视频检测主链路，只通过扩展 stats 消息和平台标定服务新增能力，因此对现有 Dashboard、GIS、告警链路影响较小。后续重点应放在标注编辑体验、annotation 版本化和生产级存储迁移。"
    )

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
    print(DOCX_PATH)
